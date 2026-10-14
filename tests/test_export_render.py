"""Export: every renderer produces a valid file, and every file carries its marking.

The classification and integrity assertions are the point of this file. They are checked
against the produced bytes, not against the code that is supposed to write them. Helpers are
disabled here; tests/test_export_helpers.py covers them with mocks. Nothing touches the network.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pytest

from vantage.export.api import ExportOptions, ExportResult, export
from vantage.models.config import ExportHelpersCfg
from vantage.models.events import RunEvent

FIXTURE = Path(__file__).parent / "fixtures" / "export" / "northwind-2026-10"
MARKING = "Restricted — Test Marking"
OFF = ExportHelpersCfg(icons=False, brand=False, stock_images=False, diagrams="none",
                       illustrations="none")


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    out = tmp_path_factory.mktemp("exports")
    events: list[RunEvent] = []
    opts = ExportOptions(dossier=FIXTURE, formats=["pptx", "docx", "html"], out_dir=out,
                         classification=MARKING, helpers=OFF, client_logo=False)
    result = export(opts, emit=events.append)
    return result, events


def _artifact(result: ExportResult, fmt: str) -> Path:
    return Path(next(a["path"] for a in result.artifacts if a["format"] == fmt))


def test_every_requested_collateral_exists(run):
    result, _ = run
    assert [a["format"] for a in result.artifacts] == ["pptx", "docx", "html"]
    for item in result.artifacts:
        assert Path(item["path"]).exists() and item["bytes"] > 1000
    assert result.manifest.exists()


def test_events_follow_the_contract(run):
    _, events = run
    types = [e.type for e in events]
    assert types[0] == "run_start" and types[-1] == "run_done"
    stages = [e.stage for e in events if e.type == "stage_start"]
    assert stages == ["ingest", "plan", "assets", "charts", "render:pptx", "render:docx",
                      "render:html"]
    artifacts = [e for e in events if e.type == "artifact"]
    assert {e.data["format"] for e in artifacts} == {"pptx", "docx", "html"}
    assert all({"format", "path", "bytes", "pages", "slides"} <= set(e.data) for e in artifacts)
    json.dumps(events[-1].data)       # run_done payload is JSON-safe


def test_result_to_dict_is_json_safe(run):
    result, _ = run
    data = json.loads(json.dumps(result.to_dict()))
    assert set(data) == {"out_dir", "manifest", "artifacts", "warnings", "plans", "verification"}
    assert data["plans"] == {"deck": "derived", "doc": "derived", "site": "derived"}
    assert data["verification"]["result"] == "PASS WITH DISCLOSURES"


def test_pptx_opens_and_marks_every_slide(run):
    from pptx import Presentation

    result, _ = run
    path = _artifact(result, "pptx")
    deck = Presentation(str(path))
    slides = list(deck.slides)
    assert len(slides) == next(a["slides"] for a in result.artifacts if a["format"] == "pptx")
    for slide in slides:
        text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        assert MARKING in text
    assert deck.core_properties.category == MARKING


def test_pptx_cover_carries_confidence_and_verification(run):
    from pptx import Presentation

    result, _ = run
    cover = Presentation(str(_artifact(result, "pptx"))).slides[0]
    text = " ".join(s.text_frame.text for s in cover.shapes if s.has_text_frame)
    assert "EQI 66.2" in text and "PASS WITH DISCLOSURES" in text and "2 blockers" in text


def test_pptx_never_prints_a_suppressed_composite_as_a_number(run):
    from pptx import Presentation

    result, _ = run
    deck = Presentation(str(_artifact(result, "pptx")))
    blob = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                blob.append(shape.text_frame.text)
            if shape.has_table:
                blob.extend(c.text for row in shape.table.rows for c in row.cells)
    text = "\n".join(blob)
    assert "SUPPRESSED" in text
    assert "O-03" not in " ".join(t for t in blob if "Model 4 attractiveness" in t)
    assert "[INTERNAL REQUIRED]" in text and "[ESTIMATED]" in text
    assert "−118" in text                                     # sign preserved


def test_docx_opens_and_marks_furniture(run):
    from docx import Document

    result, _ = run
    path = _artifact(result, "docx")
    doc = Document(str(path))
    assert len(doc.sections) == 3
    footers = " ".join(p.text for s in doc.sections for part in (s.footer, s.first_page_footer)
                       for p in part.paragraphs)
    assert MARKING in footers and "{{CLASSIFICATION}}" not in footers
    headers = " ".join(p.text for s in doc.sections for p in s.header.paragraphs)
    assert "Northwind Ceramics Corporation" in headers and "{{DOC_TITLE}}" not in headers
    assert doc.core_properties.category == MARKING


def test_docx_keeps_markers_errata_and_numbering(run):
    result, _ = run
    with zipfile.ZipFile(_artifact(result, "docx")) as zf:
        xml = zf.read("word/document.xml").decode("utf8")
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    text = re.sub(r"<[^>]+>", "", xml)
    for marker in ("[DATA UNAVAILABLE]", "[ESTIMATED]", "[INTERNAL REQUIRED]", "[DATA GAP]",
                   "SUPPRESSED", "NOT ASSESSED"):
        assert marker in text, marker
    assert "Verification erratum E1" in text
    assert "Read before use" in text and "EQI 66.2" in text
    assert "21.11" in text                     # selected appendix carried with its own number
    assert media, "charts are embedded as images"


def test_html_is_one_self_contained_file(run):
    result, _ = run
    path = _artifact(result, "html")
    page = path.read_text(encoding="utf8")
    assert not re.search(r"""(?:src|href)\s*=\s*["']?(?:https?:)?//""", page)
    assert "@import" not in page and "<link" not in page
    assert not re.search(r"""<script[^>]+src=""", page)
    assert "<svg" in page                      # charts are inline SVG
    assert page.count(MARKING) >= 3
    assert f'<meta name="classification" content="{MARKING}">' in page
    assert "marking-bar" in page
    assert "SUPPRESSED" in page and "[DATA GAP]" in page and "EQI 66.2" in page
    assert list(path.parent.glob("*.html")) == [path]


def test_svg_ids_are_unique_across_inline_charts(run):
    result, _ = run
    page = _artifact(result, "html").read_text(encoding="utf8")
    ids = re.findall(r'\sid="([^"]+)"', page)
    dupes = {i for i in ids if ids.count(i) > 1}
    assert not dupes, sorted(dupes)[:5]


def test_manifest_records_provenance(run):
    result, _ = run
    manifest = json.loads(result.manifest.read_text())
    assert manifest["classification"] == MARKING
    assert manifest["eqi"] == 66.2
    assert manifest["source_dossier"].endswith("northwind-2026-10-vantage.md")
    assert {a["format"] for a in manifest["artifacts"]} == {"pptx", "docx", "html"}


def test_html_deck_mode(tmp_path):
    result = export(ExportOptions(dossier=FIXTURE, formats=["html"], out_dir=tmp_path,
                                  site_mode="deck", depth="exec", helpers=OFF,
                                  client_logo=False))
    item = result.artifacts[0]
    page = Path(item["path"]).read_text(encoding="utf8")
    assert Path(item["path"]).name.endswith("-deck.html")
    assert page.count('class="slide') >= item["slides"] >= 10
    assert "Confidential" in page


def test_authored_plan_is_used_and_distinguished(tmp_path):
    from vantage.export.dossier import load_dossier
    from vantage.export.models import DeckPlan
    from vantage.export.pipeline import AuthoredPlanExists, write_plans

    dossier = load_dossier(FIXTURE / "northwind-2026-10-vantage.md")
    write_plans(dossier, tmp_path)
    opts = ExportOptions(dossier=FIXTURE, formats=["html"], out_dir=tmp_path, helpers=OFF,
                         client_logo=False, site_mode="deck")
    assert export(opts).plans == {"deck": "scaffold", "doc": "scaffold", "site": "scaffold"}
    authored = DeckPlan(client="Northwind", title="Northwind",
                        slides=[{"layout": "title", "title": "Hand written"}])
    (tmp_path / "plan" / "deck.json").write_text(authored.model_dump_json(), encoding="utf8")
    result = export(opts)
    assert result.plans["deck"] == "authored"
    assert "Hand written" in Path(result.artifacts[0]["path"]).read_text(encoding="utf8")
    with pytest.raises(AuthoredPlanExists):
        write_plans(dossier, tmp_path)


def test_malformed_plan_degrades_without_failing(tmp_path):
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan" / "deck.json").write_text('{"slides": [{"layout": "nope"}]}')
    result = export(ExportOptions(dossier=FIXTURE, formats=["pptx"], out_dir=tmp_path,
                                  helpers=OFF, client_logo=False))
    assert result.plans["deck"] == "derived"
    assert any("deck.json" in w for w in result.warnings)
    assert result.artifacts[0]["format"] == "pptx"


def test_unknown_format_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="unknown format"):
        export(ExportOptions(dossier=FIXTURE, formats=["keynote"], out_dir=tmp_path))


def test_find_dossier(tmp_path):
    from vantage.export.api import find_dossier

    assert find_dossier(FIXTURE).name == "northwind-2026-10-vantage.md"
    assert find_dossier(FIXTURE / "northwind-2026-10-vantage.md").is_file()
    with pytest.raises(FileNotFoundError):
        find_dossier(tmp_path)
    with pytest.raises(FileNotFoundError):
        find_dossier(tmp_path / "missing")


def test_pdf_without_libreoffice_still_builds_the_docx(tmp_path, monkeypatch):
    import vantage.export.render.pdf_render as pdf
    import vantage.export.render.whitepaper as wp

    monkeypatch.setattr(wp, "soffice_binary", lambda *_a, **_k: None)
    monkeypatch.setattr(pdf, "soffice_binary", lambda *_a, **_k: None)
    result = export(ExportOptions(dossier=FIXTURE, formats=["pdf"], out_dir=tmp_path,
                                  helpers=OFF, client_logo=False))
    assert result.artifacts == []                       # pdf not produced, docx not requested
    assert (tmp_path / "northwind-2026-10-whitepaper.docx").exists()
    assert any("PDF not produced" in w for w in result.warnings)

