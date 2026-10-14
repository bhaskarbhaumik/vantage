"""Low-level OOXML helpers python-docx does not expose.

Everything here manipulates the underlying XML directly. Kept in one module so
the renderer stays readable and the XML incantations are documented once.
"""

from __future__ import annotations

from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.shared import Pt
from docx.table import _Cell
from docx.text.paragraph import Paragraph


def _sub(parent, tag: str, **attrs):
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn(f"w:{key}"), str(value))
    parent.append(element)
    return element


def shade(cell: _Cell, hex_fill: str) -> None:
    """Solid background fill for a table cell."""
    _sub(cell._tc.get_or_add_tcPr(), "w:shd", val="clear", color="auto", fill=hex_fill.lstrip("#"))


def cell_borders(
    cell: _Cell,
    *,
    top: tuple[str, int] | None = None,
    bottom: tuple[str, int] | None = None,
    left: tuple[str, int] | None = None,
    right: tuple[str, int] | None = None,
) -> None:
    """Per-edge borders as ``(hex_colour, eighths_of_a_point)``; ``None`` clears the edge."""
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        element = OxmlElement(f"w:{edge}")
        if spec is None:
            element.set(qn("w:val"), "nil")
        else:
            colour, size = spec
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(size))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), colour.lstrip("#"))
        borders.append(element)

    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:tcBorders")):
        tc_pr.remove(existing)
    tc_pr.append(borders)


def cell_margins(cell: _Cell, top=0, start=0, bottom=0, end=0) -> None:
    """Internal padding in twips (1/20 pt)."""
    margins = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    cell._tc.get_or_add_tcPr().append(margins)


def cell_valign(cell: _Cell, align: str = "center") -> None:
    _sub(cell._tc.get_or_add_tcPr(), "w:vAlign", val=align)


def prevent_row_split(row) -> None:
    """Keep a table row on one page — stops a cell splitting mid-sentence."""
    _sub(row._tr.get_or_add_trPr(), "w:cantSplit")


def repeat_header_row(row) -> None:
    """Repeat this row at the top of every page the table spans."""
    _sub(row._tr.get_or_add_trPr(), "w:tblHeader", val="true")


def keep_with_next(paragraph: Paragraph, on: bool = True) -> None:
    """Stop a heading or lead-in stranding at the foot of a page."""
    _sub(paragraph._p.get_or_add_pPr(), "w:keepNext", val="true" if on else "false")


def paragraph_shading(paragraph: Paragraph, hex_fill: str) -> None:
    _sub(paragraph._p.get_or_add_pPr(), "w:shd", val="clear", color="auto", fill=hex_fill.lstrip("#"))


def paragraph_border(
    paragraph: Paragraph,
    edge: str,
    hex_colour: str,
    size: int = 24,
    space: int = 8,
) -> None:
    """A single rule on one edge of a paragraph — the vertical accent bar."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), hex_colour.lstrip("#"))
    borders.append(element)


def table_layout_fixed(table) -> None:
    """Honour explicit column widths instead of autofitting to content."""
    _sub(table._tbl.tblPr, "w:tblLayout", type="fixed")


def table_indent(table, twips: int) -> None:
    _sub(table._tbl.tblPr, "w:tblInd", w=twips, type="dxa")


def add_field(paragraph: Paragraph, instruction: str, placeholder: str = "") -> None:
    """Insert a Word field (``TOC``, ``PAGE``, ``NUMPAGES``, …).

    Written as the explicit begin/instrText/separate/result/end run sequence so
    both Word and LibreOffice recalculate it. ``placeholder`` is what a consumer
    that refuses to calculate will show.
    """
    run = paragraph.add_run()
    _sub(run._r, "w:fldChar", fldCharType="begin")

    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)

    run = paragraph.add_run()
    _sub(run._r, "w:fldChar", fldCharType="separate")

    if placeholder:
        paragraph.add_run(placeholder)

    run = paragraph.add_run()
    _sub(run._r, "w:fldChar", fldCharType="end")


def add_bookmark(paragraph: Paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(paragraph: Paragraph, anchor: str, text: str, colour: str, size_pt: float | None = None):
    """A clickable cross-reference to a bookmark in the same document."""
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    colour_el = OxmlElement("w:color")
    colour_el.set(qn("w:val"), colour.lstrip("#"))
    run_pr.append(colour_el)
    if size_pt:
        for tag in ("w:sz", "w:szCs"):
            sz = OxmlElement(tag)
            sz.set(qn("w:val"), str(int(size_pt * 2)))
            run_pr.append(sz)
    run.append(run_pr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)
    return link


def set_cell_width(cell: _Cell, twips: int) -> None:
    """Set the cell width, replacing any existing ``w:tcW``.

    Appending a second ``w:tcW`` is schema-invalid and silently ignored by both
    Word and LibreOffice, which then fall back to equal column widths.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:tcW")):
        tc_pr.remove(existing)
    width = OxmlElement("w:tcW")
    width.set(qn("w:w"), str(twips))
    width.set(qn("w:type"), "dxa")
    tc_pr.insert(0, width)


def style_run(run, *, font: str | None = None, size: float | None = None, colour: str | None = None,
              bold: bool | None = None, italic: bool | None = None, strike: bool | None = None,
              all_caps: bool = False, spacing: int | None = None) -> None:
    """Apply character formatting, including the east-asian/complex-script names.

    python-docx only sets ``w:ascii``/``w:hAnsi``; without ``w:cs`` a run can
    silently fall back to a different face for anything the primary font lacks.
    """
    if font:
        run.font.name = font
        r_pr = run._r.get_or_add_rPr()
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), font)
    if size is not None:
        run.font.size = Pt(size)
    if colour:
        from docx.shared import RGBColor

        run.font.color.rgb = RGBColor.from_string(colour.lstrip("#").upper())
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if strike is not None:
        run.font.strike = strike
    if all_caps:
        run.font.all_caps = True
    if spacing is not None:
        _sub(run._r.get_or_add_rPr(), "w:spacing", val=spacing)
