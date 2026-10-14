"""Render a :class:`DeckPlan` as a 16:9 PowerPoint deck.

Every layout is drawn from primitives rather than pulled from a .potx master, so the deck is
self-contained and the geometry is auditable in one file. Colour and type come from the theme.

Two things are applied centrally in ``_chrome`` so no layout can omit them: the **classification
marking** on every slide, and the provenance footer. The cover always carries the dossier's
Overall Confidence / EQI and verification result — that is the renderer's job, not the plan's, so
an authored plan cannot drop it.
"""

from __future__ import annotations

from pathlib import Path

from ..context import RenderContext
from ..models import DeckPlan, Slide, SlideLayout, Table
from ..richtext import parse_runs, plain, truncate
from ..theme import HANDLING, SCORE_DISCLAIMER


def _pptx():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    return Presentation, RGBColor, MSO_SHAPE, MSO_ANCHOR, PP_ALIGN, Emu, Inches, Pt


class PptxRenderer:
    def __init__(self, plan: DeckPlan, ctx: RenderContext, meta: dict | None = None) -> None:
        (Presentation, self.RGBColor, self.MSO_SHAPE, self.MSO_ANCHOR, self.PP_ALIGN,
         self.Emu, self.Inches, self.Pt) = _pptx()
        self.plan = plan
        self.ctx = ctx
        self.meta = meta or {}
        self.P = ctx.theme.palette
        self.T = ctx.theme.type
        self.prs = Presentation()
        In = self.Inches
        self.prs.slide_width = In(13.333)
        self.prs.slide_height = In(7.5)
        self.W, self.H = self.prs.slide_width, self.prs.slide_height
        self.MARGIN = In(0.72)
        self.CONTENT_W = self.W - self.MARGIN * 2
        self.HEADER_Y = In(0.42)
        self.TITLE_Y = In(1.12)
        self.FOOTER_Y = self.H - In(0.5)
        self._blank = self.prs.slide_layouts[6]
        self._index = 0

    # ── public API ───────────────────────────────────────────────────────────
    def render(self, out_path: Path) -> Path:
        for slide_plan in self.plan.slides:
            self._index += 1
            self._render_slide(slide_plan)
        self._set_core_properties()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out_path))
        return out_path

    def _render_slide(self, plan: Slide):
        slide = self.prs.slides.add_slide(self._blank)
        dark = plan.emphasis == "dark"
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(self.P.ink if dark else self.P.paper)
        handler = {
            SlideLayout.TITLE: self._layout_title,
            SlideLayout.SECTION: self._layout_section,
            SlideLayout.STATEMENT: self._layout_statement,
            SlideLayout.BULLETS: self._layout_bullets,
            SlideLayout.TWO_COLUMN: self._layout_two_column,
            SlideLayout.KPI: self._layout_kpi,
            SlideLayout.TABLE: self._layout_table,
            SlideLayout.CHART: self._layout_chart,
            SlideLayout.FIGURE: self._layout_figure,
            SlideLayout.QUOTE: self._layout_quote,
            SlideLayout.TIMELINE: self._layout_timeline,
            SlideLayout.MATRIX: self._layout_table,
            SlideLayout.CLOSING: self._layout_closing,
        }.get(plan.layout, self._layout_bullets)
        handler(slide, plan, dark)
        self._chrome(slide, plan, dark)
        if plan.notes:
            slide.notes_slide.notes_text_frame.text = plain(plan.notes)
        return slide

    # ── primitives ───────────────────────────────────────────────────────────
    def _rgb(self, hex_colour: str):
        return self.RGBColor.from_string(hex_colour.lstrip("#").upper())

    def _faint(self, dark: bool) -> str:
        return self.P.on_dark_faint if dark else self.P.ink_faint

    def _rect(self, slide, left, top, width, height, colour: str, shape=None, radius=None):
        shp = slide.shapes.add_shape(shape or self.MSO_SHAPE.RECTANGLE, left, top, width, height)
        shp.fill.solid()
        shp.fill.fore_color.rgb = self._rgb(colour)
        shp.line.fill.background()
        shp.shadow.inherit = False
        if radius is not None:
            shp.adjustments[0] = radius
        return shp

    def _text(self, slide, left, top, width, height, text: str, *, size: float = 14,
              colour: str | None = None, font: str | None = None, bold: bool = False,
              italic: bool = False, align=None, anchor=None, line_spacing: float = 1.12,
              markup: bool = False):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor or self.MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        para = frame.paragraphs[0]
        para.alignment = align or self.PP_ALIGN.LEFT
        para.line_spacing = line_spacing
        pieces = parse_runs(text) if markup else [type("R", (), {
            "text": plain(text), "bold": False, "italic": False, "code": False})()]
        for piece in pieces:
            run = para.add_run()
            run.text = piece.text
            self._style(run, size=size, colour=colour, font=font, bold=bold or piece.bold,
                        italic=italic or piece.italic, code=piece.code)
        return box

    def _style(self, run, *, size: float, colour: str | None, font: str | None = None,
               bold: bool = False, italic: bool = False, code: bool = False) -> None:
        run.font.size = self.Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = self.T.mono if code else (font or self.T.body)
        run.font.color.rgb = self._rgb(colour or self.P.ink)

    def _picture(self, slide, path: Path, left, top, width=None, height=None, fill_box=False):
        """Place an image; with ``fill_box`` crop it to cover ``width``×``height`` exactly."""
        try:
            pic = slide.shapes.add_picture(str(path), left, top, width=width, height=height)
        except Exception:
            return None
        if fill_box and width and height:
            try:
                from PIL import Image

                with Image.open(path) as img:
                    iw, ih = img.size
                box_ratio, img_ratio = width / height, iw / ih
                if img_ratio > box_ratio:
                    excess = 1 - box_ratio / img_ratio
                    pic.crop_left = pic.crop_right = excess / 2
                else:
                    excess = 1 - img_ratio / box_ratio
                    pic.crop_top = pic.crop_bottom = excess / 2
                pic.left, pic.top, pic.width, pic.height = left, top, width, height
            except Exception:
                pass
        return pic

    def _icon(self, slide, name: str | None, left, top, size, colour: str) -> None:
        """A get-icon glyph; without one, a small accent marker so the grid stays consistent."""
        path = self.ctx.icon_png(name, colour) if name else None
        if path is not None:
            self._picture(slide, path, left, top, width=size, height=size)
            return
        if name:
            dot = self._rect(slide, left + size * 0.3, top + size * 0.3, size * 0.4, size * 0.4,
                             colour, shape=self.MSO_SHAPE.OVAL)
            dot.name = f"icon:{name}"

    # ── chrome shared by every slide ─────────────────────────────────────────
    def _chrome(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        faint = self._faint(dark)
        # Classification — mandatory, bottom-left, every slide.
        self._text(slide, self.MARGIN, self.FOOTER_Y, In(4.2), In(0.28),
                   self.ctx.classification, size=9, colour=faint, bold=True)
        self._text(slide, self.W / 2 - In(0.5), self.FOOTER_Y, In(1.0), In(0.28),
                   str(self._index), size=9, colour=faint, align=self.PP_ALIGN.CENTER)
        if plan.source_sections:
            refs = " · ".join(f"§{s}" for s in plan.source_sections[:4])
            self._text(slide, self.W - self.MARGIN - In(4.0), self.FOOTER_Y, In(4.0), In(0.28),
                       refs, size=9, colour=faint, align=self.PP_ALIGN.RIGHT)
        logo = self.ctx.brand_logo_dark if dark else self.ctx.brand_logo
        if logo is not None and plan.layout != SlideLayout.TITLE:
            self._picture(slide, logo, self.W - self.MARGIN - In(1.3), self.HEADER_Y,
                          height=In(0.3))
        if self._has_score(plan):
            self._text(slide, self.MARGIN, self.FOOTER_Y - In(0.62), self.CONTENT_W, In(0.3),
                       SCORE_DISCLAIMER, size=8, colour=faint, italic=True)

    @staticmethod
    def _has_score(plan: Slide) -> bool:
        notes = plan.notes or ""
        chart_note = (plan.chart.note or "") if plan.chart else ""
        return SCORE_DISCLAIMER in notes or SCORE_DISCLAIMER in chart_note

    def _accent_rule(self, slide, top, height, dark: bool = False, left=None):
        bar = self._rect(slide, self.MARGIN if left is None else left, top, self.Inches(0.055),
                         height, self.P.accent_on_dark if dark else self.P.accent,
                         shape=self.MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        return bar

    def _title_block(self, slide, plan: Slide, dark: bool):
        In = self.Inches
        y = self.TITLE_Y
        left = self.MARGIN
        if plan.icon:
            self._icon(slide, plan.icon, self.MARGIN, y + In(0.04), In(0.46),
                       self.P.accent_on_dark if dark else self.P.accent)
            left = self.MARGIN + In(0.64)
        title = plain(plan.title)
        size = 28 if len(title) < 60 else 23
        self._text(slide, left, y, self.W - self.MARGIN - left, In(0.9), title, size=size,
                   colour=self.P.paper if dark else self.P.ink, font=self.T.display,
                   line_spacing=1.02)
        y = y + In(0.62 if len(title) < 60 else 1.0)
        if plan.subtitle:
            self._text(slide, left, y, self.W - self.MARGIN - left, In(0.44), plan.subtitle,
                       size=14, colour=self.P.accent_on_dark if dark else self.P.accent,
                       font=self.T.display)
            y = y + In(0.44)
        return y + In(0.18)

    # ── layouts ──────────────────────────────────────────────────────────────
    def _layout_title(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        image = self.ctx.stock(plan.image_query)
        text_w = self.CONTENT_W
        if image is not None:
            left = self.W * 0.56
            self._picture(slide, image, left, 0, width=self.W - left, height=self.H,
                          fill_box=True)
            text_w = left - self.MARGIN - In(0.4)
        top = In(1.0)
        logo = self.ctx.client_logo
        if logo is not None:
            plate = self._rect(slide, self.MARGIN, top - In(0.1), In(2.3), In(0.82), self.P.paper,
                               shape=self.MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
            plate.name = "client-logo-plate"
            self._picture(slide, logo, self.MARGIN + In(0.15), top, height=In(0.62))
        brand = self.ctx.brand_logo_dark if dark else self.ctx.brand_logo
        if brand is not None:
            self._picture(slide, brand, self.MARGIN, self.H - In(1.55), height=In(0.36))
        self._rect(slide, self.MARGIN, In(2.42), In(1.2), In(0.06),
                   self.P.accent_on_dark if dark else self.P.accent)
        self._text(slide, self.MARGIN, In(2.68), text_w, In(1.5), plan.title, size=44,
                   colour=self.P.paper if dark else self.P.ink, font=self.T.display,
                   line_spacing=1.02)
        self._text(slide, self.MARGIN, In(4.08), text_w, In(0.5), plan.subtitle, size=20,
                   colour=self.P.accent_on_dark if dark else self.P.accent, font=self.T.display)
        if plan.body:
            self._text(slide, self.MARGIN, In(4.62), text_w, In(0.4), plan.body, size=12,
                       colour=self._faint(dark))
        reliability = self._reliability_line()
        if reliability:
            self._text(slide, self.MARGIN, In(5.02), text_w, In(0.7), reliability, size=12,
                       colour=self.P.on_dark_soft if dark else self.P.ink_soft, markup=True,
                       line_spacing=1.25)

    def _reliability_line(self) -> str:
        m = self.meta
        parts = []
        if m.get("overall_confidence") or m.get("eqi") is not None:
            eqi = f" (EQI {m['eqi']:g})" if m.get("eqi") is not None else ""
            parts.append(f"**Overall confidence:** {m.get('overall_confidence') or '—'}{eqi}")
        if m.get("result"):
            counts = " · ".join(f"{m[k]} {k}" for k in ("blockers", "majors", "minors")
                                if m.get(k) is not None)
            parts.append(f"**Verification:** {m['result']}" + (f" ({counts})" if counts else ""))
        return "\n".join(parts).replace("\n", " · ")

    def _layout_section(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        image = self.ctx.stock(plan.image_query)
        text_w = self.CONTENT_W - In(0.34)
        if image is not None:
            left = self.W * 0.62
            self._picture(slide, image, left, 0, width=self.W - left, height=self.H,
                          fill_box=True)
            text_w = left - self.MARGIN - In(0.7)
        if plan.icon:
            self._icon(slide, plan.icon, self.MARGIN + In(0.34), In(2.1), In(0.7),
                       self.P.accent_on_dark if dark else self.P.accent)
        self._accent_rule(slide, In(3.05), In(1.2), dark)
        self._text(slide, self.MARGIN + In(0.34), In(3.0), text_w, In(0.9), plan.title,
                   size=38, colour=self.P.paper if dark else self.P.ink, font=self.T.display)
        if plan.subtitle:
            self._text(slide, self.MARGIN + In(0.34), In(3.86), text_w, In(0.5), plan.subtitle,
                       size=16, colour=self.P.accent_on_dark if dark else self.P.accent,
                       font=self.T.display)

    def _layout_statement(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        if plan.icon:
            self._icon(slide, plan.icon, self.MARGIN + In(0.34), In(1.05), In(0.5),
                       self.P.accent_on_dark if dark else self.P.accent)
        title = plain(plan.title)
        size = 34 if len(title) < 90 else 27 if len(title) < 150 else 22
        self._accent_rule(slide, In(1.9), In(2.6), dark)
        self._text(slide, self.MARGIN + In(0.34), In(1.85), self.CONTENT_W - In(0.9), In(2.4),
                   plan.title, size=size, colour=self.P.paper if dark else self.P.ink,
                   font=self.T.display, line_spacing=1.14, markup=True)
        if plan.subtitle:
            self._text(slide, self.MARGIN + In(0.34), In(4.28), self.CONTENT_W - In(0.9),
                       In(0.36), plan.subtitle, size=11,
                       colour=self.P.accent_on_dark if dark else self.P.accent)
        if plan.body:
            self._text(slide, self.MARGIN + In(0.34), In(4.72), self.CONTENT_W - In(0.9),
                       In(1.5), plan.body, size=14,
                       colour=self.P.on_dark_soft if dark else self.P.ink_soft,
                       line_spacing=1.3, markup=True)
        self._claim_refs(slide, plan, dark)

    def _layout_bullets(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        y = self._title_block(slide, plan, dark)
        if plan.body:
            self._text(slide, self.MARGIN, y, self.CONTENT_W, In(0.8), plan.body, size=14,
                       colour=self.P.on_dark_soft if dark else self.P.ink_soft,
                       line_spacing=1.3, markup=True)
            y = y + In(0.8)
        if plan.bullets:
            self._bullet_list(slide, self.MARGIN, y, self.CONTENT_W, plan.bullets, dark)
        self._claim_refs(slide, plan, dark)

    def _bullet_list(self, slide, left, top, width, bullets: list[str], dark: bool,
                     size: float = 13.5) -> None:
        available = float(self.Emu(self.FOOTER_Y - top).inches) - 0.45
        box = slide.shapes.add_textbox(left, top, width, self.Inches(max(available, 0.6)))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        count = len(bullets[:7])
        if count >= 6:
            size -= 1.5
        for index, bullet in enumerate(bullets[:7]):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.line_spacing = 1.2
            para.space_after = self.Pt(8)
            marker = para.add_run()
            marker.text = "— "
            self._style(marker, size=size, colour=self.P.accent_on_dark if dark else self.P.accent,
                        bold=True)
            for piece in parse_runs(truncate(bullet, 230)):
                run = para.add_run()
                run.text = piece.text
                self._style(run, size=size, colour=self.P.paper if dark else self.P.ink,
                            bold=piece.bold, italic=piece.italic, code=piece.code)

    def _layout_two_column(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        y = self._title_block(slide, plan, dark)
        gutter = In(0.5)
        col_w = (self.CONTENT_W - gutter) / 2
        for index, (heading, bullets) in enumerate(
                zip(plan.column_headings or ["", ""], [plan.bullets, plan.bullets_right],
                    strict=False)):
            left = self.MARGIN + (col_w + gutter) * index
            top = y
            if heading:
                self._text(slide, left, top, col_w, In(0.34), heading, size=12, bold=True,
                           colour=self.P.accent_on_dark if dark else self.P.accent,
                           font=self.T.display)
                top = top + In(0.42)
            if bullets:
                self._bullet_list(slide, left, top, col_w, bullets, dark, size=12)
        self._claim_refs(slide, plan, dark)

    def _layout_kpi(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        y = self._title_block(slide, plan, dark)
        kpis = plan.kpis[:4]
        if not kpis:
            return
        gutter = In(0.34)
        tile_w = (self.CONTENT_W - gutter * (len(kpis) - 1)) / len(kpis)
        tile_h = In(2.3)
        for index, kpi in enumerate(kpis):
            left = self.MARGIN + (tile_w + gutter) * index
            self._rect(slide, left, y, tile_w, tile_h,
                       self.P.dark_surface if dark else self.P.surface,
                       shape=self.MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
            accent = {"good": self.P.green, "risk": self.P.red}.get(kpi.tone, self.P.accent)
            self._rect(slide, left + In(0.24), y + In(0.3), In(0.05), In(0.5), accent,
                       shape=self.MSO_SHAPE.ROUNDED_RECTANGLE)
            if kpi.icon:
                self._icon(slide, kpi.icon, left + tile_w - In(0.62), y + In(0.24), In(0.36),
                           accent)
            inner, inner_w = left + In(0.44), tile_w - In(0.68)
            value = plain(kpi.value)
            self._text(slide, inner, y + In(0.26), inner_w - In(0.3), In(0.66), value,
                       size=30 if len(value) <= 8 else 22,
                       colour=self.P.paper if dark else self.P.ink, font=self.T.display,
                       line_spacing=1.0)
            self._text(slide, inner, y + In(1.0), inner_w, In(0.8), kpi.label, size=11,
                       colour=self.P.on_dark_soft if dark else self.P.ink_soft, line_spacing=1.2)
            if kpi.delta:
                self._text(slide, inner, y + In(1.66), inner_w, In(0.5), truncate(kpi.delta, 70),
                           size=10, colour=accent, bold=True)
        self._claim_refs(slide, plan, dark)

    def _layout_table(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        y = self._title_block(slide, plan, dark)
        table = plan.table
        if table is None or not (table.header or table.rows):
            return
        cols = max(len(table.header), max((len(r) for r in table.rows), default=0))
        rows = table.rows[:8]
        room = float(self.Emu(self.FOOTER_Y - y).inches) - 0.9
        shape = slide.shapes.add_table(len(rows) + 1, cols, self.MARGIN, y, self.CONTENT_W,
                                       In(min(0.42 * (len(rows) + 1), max(room, 1.0)))).table
        for index, width in enumerate(self._column_widths(table, cols)):
            shape.columns[index].width = self.Emu(int(width))
        body_pt = 10.5 if cols <= 4 else 9 if cols <= 6 else 8
        long_cells = sum(len(plain(c)) for r in rows for c in r) / max(1, len(rows) * cols)
        if long_cells > 90:
            body_pt -= 1.5
        for index in range(cols):
            self._cell(shape.cell(0, index), table.header[index] if index < len(table.header)
                       else "", size=body_pt, colour=self.P.paper, fill=self.P.accent, bold=True,
                       font=self.T.display)
        for row_index, row in enumerate(rows, start=1):
            for col_index in range(cols):
                text = row[col_index] if col_index < len(row) else ""
                zebra = row_index % 2 == 0
                self._cell(shape.cell(row_index, col_index), text, size=body_pt,
                           colour=self.P.paper if dark else self.P.ink,
                           fill=(self.P.dark_surface if dark else self.P.surface) if zebra
                           else (self.P.ink if dark else self.P.paper),
                           limit=180 if long_cells < 90 else 140)
        if table.caption:
            self._text(slide, self.MARGIN, self.FOOTER_Y - In(0.34), self.CONTENT_W - In(4.2),
                       In(0.3), table.caption, size=9, colour=self._faint(dark), italic=True)
        if plan.body and plan.layout == SlideLayout.MATRIX:
            slide.notes_slide.notes_text_frame.text = plain(plan.body)

    MIN_COL_IN = 0.62

    def _column_widths(self, table: Table, cols: int) -> list[float]:
        weights: list[float] = []
        for col in range(cols):
            samples = sorted(len(plain(r[col])) for r in table.rows if col < len(r))
            head = len(plain(table.header[col])) if col < len(table.header) else 0
            typical = samples[int(len(samples) * 0.75)] if samples else 0
            weights.append(max(2.5, min(float(max(typical, head)), 120.0)))
        usable = float(self.Emu(self.CONTENT_W).inches)
        total = sum(weights) or 1.0
        widths = [usable * w / total for w in weights]
        floor = min(self.MIN_COL_IN, usable / cols)
        deficit = sum(floor - w for w in widths if w < floor)
        widths = [max(w, floor) for w in widths]
        if deficit > 0:
            surplus = sum(w - floor for w in widths if w > floor) or 1.0
            widths = [w - (w - floor) / surplus * deficit if w > floor else w for w in widths]
        scale = usable / sum(widths)
        return [self.Inches(w * scale) for w in widths]

    def _cell(self, cell, text: str, *, size: float, colour: str, fill: str,
              bold: bool = False, font: str | None = None, limit: int = 180) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = self._rgb(fill)
        cell.margin_left = cell.margin_right = self.Inches(0.08)
        cell.margin_top = cell.margin_bottom = self.Inches(0.04)
        cell.vertical_anchor = self.MSO_ANCHOR.MIDDLE
        frame = cell.text_frame
        frame.word_wrap = True
        para = frame.paragraphs[0]
        para.line_spacing = 1.08
        for piece in parse_runs(truncate(text, limit)):
            run = para.add_run()
            run.text = piece.text
            self._style(run, size=size, colour=colour, font=font, bold=bold or piece.bold,
                        italic=piece.italic, code=piece.code)

    def _aside(self, slide, plan: Slide, dark: bool, y, left) -> None:
        if not plan.body:
            return
        In = self.Inches
        self._accent_rule(slide, y + In(0.1), In(1.6), dark, left=left)
        left = left + In(0.26)
        box = slide.shapes.add_textbox(left, y, self.W - self.MARGIN - left, In(3.4))
        frame = box.text_frame
        frame.word_wrap = True
        para = frame.paragraphs[0]
        para.line_spacing = 1.3
        for piece in parse_runs(plan.body):
            run = para.add_run()
            run.text = piece.text
            self._style(run, size=13, colour=self.P.paper if dark else self.P.ink_soft,
                        bold=piece.bold, italic=piece.italic, code=piece.code)

    def _layout_chart(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        y = self._title_block(slide, plan, dark)
        if plan.chart is None:
            return
        image = self.ctx.charts.png(plan.chart, width_in=7.6, height_in=4.0)
        if image is None:
            return
        room = float(self.Emu(self.FOOTER_Y - y).inches) - 0.75
        pic = self._picture(slide, image, self.MARGIN, y, height=In(min(3.9, room)))
        if pic is not None and pic.width > In(8.0):
            ratio = In(8.0) / pic.width
            pic.width, pic.height = In(8.0), int(pic.height * ratio)
        self._aside(slide, plan, dark, y, self.MARGIN + In(8.3))
        if plan.chart.note and SCORE_DISCLAIMER not in plan.chart.note:
            self._text(slide, self.MARGIN, self.FOOTER_Y - In(0.36), In(8.0), In(0.32),
                       truncate(plan.chart.note, 160), size=8, colour=self._faint(dark),
                       italic=True, line_spacing=1.1)

    def _layout_figure(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        y = self._title_block(slide, plan, dark)
        if plan.figure is None:
            return
        art = self.ctx.figures.render(plan.figure)
        room_h = float(self.Emu(self.FOOTER_Y - y).inches) - 0.7
        if art is not None and art.png is not None:
            pic = self._picture(slide, art.png, self.MARGIN, y, height=In(room_h))
            if pic is not None and pic.width > In(8.6):
                ratio = In(8.6) / pic.width
                pic.width, pic.height = In(8.6), int(pic.height * ratio)
            self._aside(slide, plan, dark, y, self.MARGIN + In(8.9))
        else:
            from ..visuals import figure_table

            table = figure_table(plan.figure)
            if table is not None:
                from ..plan import trim_table

                self._layout_table(slide, plan.model_copy(update={
                    "table": trim_table(table, 8, 4), "title": "", "subtitle": "", "icon": None,
                }), dark)
        self._text(slide, self.MARGIN, self.FOOTER_Y - In(0.36), self.CONTENT_W, In(0.32),
                   truncate(plan.figure.caption(), 170), size=8, colour=self._faint(dark),
                   italic=True)

    def _layout_quote(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        self._accent_rule(slide, In(2.2), In(2.6), dark)
        self._text(slide, self.MARGIN + In(0.4), In(2.15), self.CONTENT_W - In(1.2), In(2.7),
                   f"“{plain(plan.title)}”", size=28,
                   colour=self.P.paper if dark else self.P.ink, font=self.T.display,
                   italic=True, line_spacing=1.2)
        if plan.attribution:
            self._text(slide, self.MARGIN + In(0.4), In(5.1), self.CONTENT_W - In(1.2), In(0.5),
                       plan.attribution, size=13,
                       colour=self.P.accent_on_dark if dark else self.P.accent)

    def _layout_timeline(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        y = self._title_block(slide, plan, dark)
        steps = plan.timeline[:5]
        if not steps:
            return
        gutter = In(0.3)
        col_w = (self.CONTENT_W - gutter * (len(steps) - 1)) / len(steps)
        track_y = y + In(0.34)
        self._rect(slide, self.MARGIN, track_y, self.CONTENT_W, In(0.02),
                   self.P.dark_surface if dark else self.P.rule)
        for index, step in enumerate(steps):
            left = self.MARGIN + (col_w + gutter) * index
            self._rect(slide, left, track_y - In(0.075), In(0.17), In(0.17), self.P.accent,
                       shape=self.MSO_SHAPE.OVAL)
            self._text(slide, left, track_y + In(0.24), col_w, In(0.4), plain(step.when),
                       size=13, colour=self.P.accent_on_dark if dark else self.P.accent,
                       font=self.T.display, bold=True)
            self._text(slide, left, track_y + In(0.68), col_w, In(2.0), step.what, size=11.5,
                       colour=self.P.paper if dark else self.P.ink, line_spacing=1.24,
                       markup=True)
            if step.detail:
                self._text(slide, left, track_y + In(2.5), col_w, In(0.6), plain(step.detail),
                           size=10, colour=self._faint(dark))
        self._claim_refs(slide, plan, dark)

    def _layout_closing(self, slide, plan: Slide, dark: bool) -> None:
        In = self.Inches
        self._text(slide, self.MARGIN, In(2.9), self.CONTENT_W, In(1.0),
                   plan.title or "Vantage", size=36,
                   colour=self.P.paper if dark else self.P.ink, font=self.T.display)
        self._text(slide, self.MARGIN, In(4.1), In(9.4), In(1.4),
                   plan.subtitle or SCORE_DISCLAIMER, size=11.5, colour=self._faint(dark),
                   italic=True, line_spacing=1.3)
        self._text(slide, self.MARGIN, In(5.7), self.CONTENT_W, In(0.34),
                   f"{self.ctx.classification} · {HANDLING}", size=11,
                   colour=self.P.accent_on_dark if dark else self.P.accent, font=self.T.display)

    def _claim_refs(self, slide, plan: Slide, dark: bool) -> None:
        if not plan.claim_refs:
            return
        offset = 0.62 if self._has_score(plan) else 0.32
        self._text(slide, self.MARGIN, self.FOOTER_Y - self.Inches(offset), self.Inches(9.0),
                   self.Inches(0.28), " · ".join(plan.claim_refs[:8]), size=8,
                   colour=self._faint(dark), font=self.T.mono)

    def _set_core_properties(self) -> None:
        props = self.prs.core_properties
        props.title = f"{plain(self.plan.client)} — {self.plan.subtitle}"
        props.subject = self.plan.subtitle
        props.author = self.ctx.theme.brand_name or "Vantage"
        props.category = self.ctx.classification
        props.comments = f"{self.ctx.classification}. {HANDLING} {SCORE_DISCLAIMER}"
        props.keywords = f"{self.ctx.classification}; Vantage dossier; internal"


def render_pptx(plan: DeckPlan, out_path: Path, ctx: RenderContext,
                meta: dict | None = None) -> Path:
    return PptxRenderer(plan, ctx, meta).render(out_path)
