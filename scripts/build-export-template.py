#!/usr/bin/env python3
"""Generate the neutral whitepaper template used by ``vantage export``.

The template is built from python-docx primitives — no third-party design is copied — and holds
only what python-docx cannot conveniently author per render: three sections with their own page
furniture, page-number fields and heading styles. The renderer fills it and re-colours the
headings from the active theme, so ``--brand`` theming applies without a second template.

    section 0  cover, then document control / about / handling   (first-page footer differs)
    section 1  table of contents
    section 2  the body — running title header; classification + "Page N" footer,
               numbering restarted at 1 (the PDF contents pass reads these numbers back)

Placeholders ``{{DOC_TITLE}}`` and ``{{CLASSIFICATION}}`` sit in single runs; the renderer
substitutes them per export.

Usage:
    uv run python scripts/build-export-template.py [-o PATH]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

DEFAULT_OUT = REPO / "src" / "vantage" / "export" / "assets" / "templates" / "vantage-whitepaper.docx"


def build(out: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml.parser import OxmlElement
    from docx.shared import Inches, Pt, RGBColor

    from vantage.export.render import ooxml as ox
    from vantage.export.theme import DEFAULT_THEME

    p, t = DEFAULT_THEME.palette, DEFAULT_THEME.type
    doc = Document()

    # ── styles ───────────────────────────────────────────────────────────────
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = t.body
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(p.ink)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, font, colour, before, after in (
        ("Title", 30, t.display, p.ink, 0, 6),
        ("Heading 1", 20, t.display, p.accent_deep, 18, 8),
        ("Heading 2", 14, t.display, p.accent_deep, 14, 6),
        ("Heading 3", 11.5, t.body, p.ink, 10, 4),
    ):
        style = styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = name == "Heading 3"
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string(colour)
        r_pr = style.element.get_or_add_rPr()
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), font)
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            fonts.attrib.pop(qn(theme_attr), None)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if name == "Title":
            # The default Title style draws a bottom border; the cover does its own rule.
            p_pr = style.element.get_or_add_pPr()
            for border in p_pr.findall(qn("w:pBdr")):
                p_pr.remove(border)
    caption = styles["Caption"]
    caption.font.name = t.body
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(p.ink_soft)

    # ── page set-up: three sections ──────────────────────────────────────────
    def page(section) -> None:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(0.9)
        section.header_distance = Inches(0.45)
        section.footer_distance = Inches(0.4)

    def furniture(paragraph, left: str, right_field: bool = False, size: float = 8.5) -> None:
        paragraph.text = ""
        width = Inches(6.5)
        paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        run = paragraph.add_run(left)
        ox.style_run(run, font=t.body, size=size, colour=p.ink_faint, bold=True)
        if right_field:
            paragraph.add_run("\t")
            # "Page N" only: SECTIONPAGES is not honoured by every consumer (LibreOffice prints
            # "of 1"), and NUMPAGES would count the cover and contents the body does not.
            run = paragraph.add_run("Page ")
            ox.style_run(run, font=t.body, size=size, colour=p.ink_faint)
            ox.add_field(paragraph, "PAGE", placeholder="1")

    cover = doc.sections[0]
    page(cover)
    cover.different_first_page_header_footer = True
    furniture(cover.first_page_footer.paragraphs[0], "{{CLASSIFICATION}}")
    cover.first_page_footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    furniture(cover.footer.paragraphs[0], "{{CLASSIFICATION}}")
    head = cover.header.paragraphs[0]
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ox.style_run(head.add_run("{{DOC_TITLE}}"), font=t.body, size=8.5, colour=p.ink_faint)

    contents = doc.add_section(WD_SECTION.NEW_PAGE)
    page(contents)
    contents.different_first_page_header_footer = False
    for part in (contents.header, contents.footer):
        part.is_linked_to_previous = False
    furniture(contents.footer.paragraphs[0], "{{CLASSIFICATION}}")
    head = contents.header.paragraphs[0]
    head.text = ""
    head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ox.style_run(head.add_run("Contents"), font=t.display, size=16, colour=p.accent_deep)
    ox.paragraph_border(head, "bottom", p.rule, size=6, space=4)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    page(body)
    body.different_first_page_header_footer = False
    for part in (body.header, body.footer, body.first_page_header, body.first_page_footer):
        part.is_linked_to_previous = False
    head = body.header.paragraphs[0]
    head.text = ""
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ox.style_run(head.add_run("{{DOC_TITLE}}"), font=t.body, size=8.5, colour=p.ink_faint)
    ox.paragraph_border(head, "bottom", p.rule, size=4, space=4)
    furniture(body.footer.paragraphs[0], "{{CLASSIFICATION}}", right_field=True)
    # Body numbering restarts at 1, so the printed page numbers the PDF contents pass reads
    # back do not move when the contents page grows.
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    body._sectPr.append(pg)

    # Drop the empty paragraph python-docx's default document starts with, so the body holds
    # only the two section-break paragraphs and the final sectPr.
    for para in list(doc.paragraphs):
        if para._p.pPr is None or para._p.pPr.find(qn("w:sectPr")) is None:
            para._p.getparent().remove(para._p)

    # Ask the consumer to recalculate fields (Page N, and the TOC field in docx-only runs).
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    props = doc.core_properties
    props.title = props.subject = props.keywords = props.comments = ""
    props.author = "Vantage"
    props.last_modified_by = "Vantage"

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-o", "--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    out = build(args.out)
    try:
        shown = out.relative_to(REPO)
    except ValueError:
        shown = out
    print(f"template → {shown}  ({out.stat().st_size / 1024:.0f} KB, 3 sections)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
